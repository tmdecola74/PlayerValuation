"""Read-only directory inventory; does not infer seasons or projection systems."""
import argparse,csv,json
from pathlib import Path
from openpyxl import load_workbook
p=argparse.ArgumentParser(description=__doc__)
p.add_argument('directories',nargs='+')
a=p.parse_args()
for directory in a.directories:
 for f in sorted(Path(directory).rglob('*')):
  if f.suffix.lower() not in ('.csv','.tsv','.xlsx') or f.name.startswith('~$'): continue
  result={'path':str(f),'bytes':f.stat().st_size}
  try:
   if f.suffix.lower()=='.xlsx':
    w=load_workbook(f,read_only=True,data_only=True)
    try:result['sheets']={s.title:[list(r) for r in s.iter_rows(max_row=4,values_only=True)] for s in w}
    finally:w.close()
   else:
    with f.open(encoding='utf-8-sig',newline='') as stream:result['columns']=next(csv.reader(stream,delimiter='\t' if f.suffix.lower()=='.tsv' else ','))
  except Exception as e:result['error']=str(e)
  print(json.dumps(result,default=str))
