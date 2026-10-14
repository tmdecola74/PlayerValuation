"""PlayerValuation warehouse foundation. Single writer; immutable file versions."""
from pathlib import Path
import argparse, csv, hashlib, json, math, re, shutil, uuid
from datetime import datetime, timezone
import duckdb
import polars as pl
import yaml
from openpyxl import load_workbook

VERSION = '1.0.0'
DATASETS = ('projection_hitter','projection_pitcher','actual_hitter','actual_pitcher')
TEXT = ('player_key','mlbam_id','fangraphs_id','name','team','source','snapshot_date','available_at','source_file','source_hash','run_id')
NUMERIC = ('PA','AB','H','R','HR','RBI','SB','CS','BB','SO','AVG','OBP','SLG','OPS','IP','ER','HA','BBA','W','SV','HLD','GS','G','ERA','WHIP','QS')
SCHEMA = {'season':pl.Int64, **{k:pl.String for k in TEXT}, **{k:pl.Float64 for k in NUMERIC}}

def digest(path):
    h=hashlib.sha256()
    with open(path,'rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''): h.update(b)
    return h.hexdigest()

def literal(s): return "'"+str(s).replace("'","''")+"'"

def load_config(path):
    path=Path(path).resolve()
    c=yaml.safe_load(path.read_text(encoding='utf-8'))
    root=Path(c.get('project_root','..'))
    c['_root']=(root if root.is_absolute() else path.parent/root).resolve()
    return c

def location(c,key,default):
    p=Path(c.get(key,default))
    return p if p.is_absolute() else c['_root']/p

def connect(c):
    db=location(c,'database','data/database/player_valuation.duckdb')
    db.parent.mkdir(parents=True,exist_ok=True)
    return duckdb.connect(str(db))

def refresh_views(con):
    for ds in DATASETS:
        paths=[r[0] for r in con.execute('SELECT parquet_path FROM dataset_version WHERE dataset=? AND active',[ds]).fetchall()]
        if paths:
            query='SELECT * FROM read_parquet(['+','.join(literal(p) for p in paths)+'], union_by_name=true, hive_partitioning=false)'
        else:
            types={pl.Int64:'BIGINT',pl.String:'VARCHAR',pl.Float64:'DOUBLE'}
            query='SELECT '+','.join(f'CAST(NULL AS {types[t]}) AS "{k}"' for k,t in SCHEMA.items())+' WHERE FALSE'
        con.execute(f'CREATE OR REPLACE VIEW {ds} AS {query}')
    con.execute('CREATE OR REPLACE VIEW dim_player AS SELECT DISTINCT player_key, mlbam_id, fangraphs_id, name FROM ('+' UNION ALL '.join('SELECT player_key,mlbam_id,fangraphs_id,name FROM '+d for d in DATASETS)+')')

def initialize(c):
    for d in ('data/raw','data/parquet','data/database','audit','output'): (c['_root']/d).mkdir(parents=True,exist_ok=True)
    with connect(c) as con:
        con.execute((Path(__file__).parents[2]/'sql/schema_v1_0.sql').read_text())
        refresh_views(con)

def read_rows(path,s):
    if path.suffix.lower() in ('.csv','.tsv'):
        with open(path,encoding=s.get('encoding','utf-8-sig'),newline='') as f:
            reader=csv.DictReader(f,delimiter=s.get('delimiter','\t' if path.suffix.lower()=='.tsv' else ','))
            headers=reader.fieldnames or []
            if len(headers)!=len(set(headers)): raise ValueError('Duplicate source column headings')
            return headers,list(reader)
    if path.suffix.lower()=='.xlsx':
        if not s.get('sheet'): raise ValueError('XLSX requires explicit sheet name')
        wb=load_workbook(path,read_only=True,data_only=True)
        try:
            ws=wb[s['sheet']]
            it=ws.iter_rows(min_row=s.get('header_row',1),values_only=True)
            headers=[str(x).strip() if x is not None else '' for x in next(it)]
            nonblank=[x for x in headers if x]
            if len(nonblank)!=len(set(nonblank)): raise ValueError('Duplicate XLSX headings')
            return headers,[dict(zip(headers,r)) for r in it if any(x is not None for x in r)]
        finally: wb.close()
    raise ValueError('Only CSV, TSV and XLSX sources supported in this builder')

def identifier(v):
    if v is None or str(v).strip()=='': return None
    v=str(v).strip()
    if re.fullmatch(r'\d+\.0',v): v=v[:-2]
    return v

def number(v,field,s):
    if v is None or str(v).strip() in ('','-','NA','N/A'): return None
    n=float(str(v).replace(',','').strip())
    if not math.isfinite(n): raise ValueError('Non-finite number')
    if field=='IP' and s.get('ip_format','decimal')=='baseball':
        whole=int(n); outs=round((n-whole)*10)
        if outs not in (0,1,2) or abs(n-(whole+outs/10))>1e-6: raise ValueError('Invalid baseball innings')
        n=whole+outs/3
    if field in ('AVG','OBP','SLG','OPS') and s.get('rate_scale','fraction')=='percent': n/=100
    if n<0: raise ValueError('Negative statistic')
    if field in ('AVG','OBP') and n>1: raise ValueError('Rate exceeds 1; check rate_scale')
    return n

def normalize(rows,headers,s,path,file_hash,run):
    mapping=s.get('columns',{})
    unknown=set(mapping)-set(SCHEMA)
    if unknown: raise ValueError(f'Unknown canonical columns: {sorted(unknown)}')
    missing=set(mapping.values())-set(headers)
    if missing: raise ValueError(f'Mapped columns missing: {sorted(missing)}')
    if 'name' not in mapping: raise ValueError('Map name explicitly')
    records=[]; issues=[]; seen=set()
    for index,row in enumerate(rows, s.get('header_row',1)+1):
        r={k:None for k in SCHEMA}
        r.update(season=int(s['season']),source=s['source'],snapshot_date=s.get('snapshot_date'),available_at=s.get('available_at'),source_file=str(path),source_hash=file_hash,run_id=run)
        for k,col in mapping.items():
            if k in NUMERIC:
                try:r[k]=number(row.get(col),k,s)
                except (ValueError,TypeError) as e: issues.append({'row':index,'code':'INVALID_NUMBER','field':k,'value':str(row.get(col)),'detail':str(e)})
            elif k in ('mlbam_id','fangraphs_id'): r[k]=identifier(row.get(col))
            elif k not in ('season','run_id','source','source_hash','source_file'): r[k]=str(row[col]).strip() if row.get(col) is not None else None
        if r['mlbam_id'] and not re.fullmatch(r'[1-9]\d*',r['mlbam_id']):
            issues.append({'row':index,'code':'INVALID_ID','field':'mlbam_id','value':r['mlbam_id'],'detail':'MLBAM must be a positive integer'})
        r['player_key']='MLBAM:'+r['mlbam_id'] if r['mlbam_id'] else ('FG:'+r['fangraphs_id'] if r['fangraphs_id'] else None)
        if not r['name'] or not r['player_key']: issues.append({'row':index,'code':'MISSING_IDENTITY','field':'player_key','value':r['name'],'detail':'Supply MLBAM or FanGraphs ID; names are never fuzzy matched'})
        if r['player_key'] in seen: issues.append({'row':index,'code':'DUPLICATE_PLAYER','field':'player_key','value':r['player_key'],'detail':'Use a season-total sheet; team stints require an explicit aggregation adapter'})
        seen.add(r['player_key'])
        records.append(r)
    return pl.DataFrame(records,schema=SCHEMA),issues

def build(c,manifest,inspect=False):
    specs=yaml.safe_load(Path(manifest).read_text(encoding='utf-8')).get('sources',[])
    if not specs: raise ValueError('No sources configured; edit config/data_sources.yaml first')
    if not inspect: initialize(c)
    result=[]
    for s in specs:
        ds=s['dataset']
        if ds not in DATASETS: raise ValueError(f'Unsupported dataset {ds}')
        p=Path(s['path']); p=p if p.is_absolute() else c['_root']/p
        headers,rows=read_rows(p,s)
        if inspect:
            result.append({'path':str(p),'rows':len(rows),'columns':headers}); continue
        run=uuid.uuid4().hex; h=digest(p)
        df,issues=normalize(rows,headers,s,p,h,run)
        audit=location(c,'audit','audit')/run; audit.mkdir(parents=True,exist_ok=True)
        (audit/'issues.json').write_text(json.dumps(issues,indent=2),encoding='utf-8')
        if issues or not len(rows): raise ValueError(f'Rejected {p}: {len(issues)} issues or empty source; see {audit}')
        fingerprint=hashlib.sha256(json.dumps({'source':s,'hash':h,'version':VERSION},sort_keys=True).encode()).hexdigest()
        with connect(c) as con:
            exists=con.execute('SELECT version_id FROM dataset_version WHERE fingerprint=? AND active',[fingerprint]).fetchone()
            if exists: result.append({'path':str(p),'status':'UNCHANGED'}); continue
            raw=location(c,'raw','data/raw')/h/p.name
            raw.parent.mkdir(parents=True,exist_ok=True)
            if not raw.exists(): shutil.copy2(p,raw)
            elif digest(raw)!=h: raise ValueError('Archived source hash mismatch')
            out=location(c,'parquet','data/parquet')/ds/f'season={s["season"]}'/run/'part.parquet'
            out.parent.mkdir(parents=True,exist_ok=True)
            tmp=out.with_suffix('.tmp'); df.write_parquet(tmp,compression='zstd'); tmp.replace(out)
            con.execute('BEGIN')
            try:
                con.execute('UPDATE dataset_version SET active=false WHERE dataset=? AND season=? AND source=?',[ds,int(s['season']),s['source']])
                con.execute('INSERT INTO dataset_version VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)',[run,ds,int(s['season']),s['source'],h,fingerprint,str(raw.resolve()),str(out.resolve()),digest(out),len(rows),len(df),datetime.now(timezone.utc).isoformat(),True])
                refresh_views(con)
                con.execute('COMMIT')
            except Exception:
                con.execute('ROLLBACK'); raise
        result.append({'path':str(p),'status':'INGESTED','rows':len(df),'run_id':run,'unmapped_mlbam':df['mlbam_id'].null_count()})
    return result

def validate(c):
    checks=[]
    with connect(c) as con:
        versions=con.execute('SELECT version_id,dataset,raw_path,parquet_path,source_hash,parquet_hash,source_rows,normalized_rows FROM dataset_version WHERE active').fetchall()
        for run,ds,raw,pq,sh,ph,sr,nr in versions:
            try:
                df=pl.read_parquet(pq)
                ok=digest(Path(raw))==sh and digest(Path(pq))==ph and len(df)==sr==nr
                checks.append({'check':'FILE_RECONCILIATION','dataset':ds,'version':run,'pass':ok,'rows':len(df)})
                checks.append({'check':'IDENTITY_UNIQUENESS','dataset':ds,'version':run,'pass':df['player_key'].null_count()==0 and df['player_key'].n_unique()==len(df)})
                checks.append({'check':'MLBAM_COVERAGE','dataset':ds,'version':run,'pass':True,'unresolved':df['mlbam_id'].null_count(),'severity':'WARNING' if df['mlbam_id'].null_count() else 'INFO'})
                for field in NUMERIC:
                    checks.append({'check':'NULL_RATE','dataset':ds,'version':run,'field':field,'nulls':df[field].null_count(),'rows':len(df),'pass':True})
            except Exception as e: checks.append({'check':'FILE_READ','dataset':ds,'version':run,'pass':False,'error':str(e)})
        for ds in DATASETS:
            expected=sum(r[7] for r in versions if r[1]==ds)
            actual=con.execute(f'SELECT count(*) FROM {ds}').fetchone()[0]
            checks.append({'check':'DUCKDB_ROW_RECONCILIATION','dataset':ds,'pass':expected==actual,'expected':expected,'actual':actual})
        checks.append({'check':'NONEMPTY_WAREHOUSE','pass':bool(versions)})
    dest=location(c,'audit','audit'); dest.mkdir(parents=True,exist_ok=True)
    report={'generated_at':datetime.now(timezone.utc).isoformat(),'script_version':VERSION,'pass':all(x['pass'] for x in checks),'checks':checks}
    (dest/'warehouse_validation.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    keys=sorted(set().union(*(r.keys() for r in checks)))
    with open(dest/'warehouse_validation.csv','w',newline='',encoding='utf-8') as f:
        writer=csv.DictWriter(f,fieldnames=keys); writer.writeheader(); writer.writerows(checks)
    return report

def main(mode):
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config',default='config/warehouse.yaml')
    parser.add_argument('--sources',default='config/data_sources.yaml')
    parser.add_argument('--inspect',action='store_true',help='Read source headings/row counts without writing warehouse')
    a=parser.parse_args(); c=load_config(a.config)
    if mode=='init': initialize(c); result={'status':'INITIALIZED'}
    elif mode=='build': result=build(c,a.sources,a.inspect)
    else:
        result=validate(c)
        print(json.dumps(result,indent=2)); return 0 if result['pass'] else 1
    print(json.dumps(result,indent=2)); return 0
